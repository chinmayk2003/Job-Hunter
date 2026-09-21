"""
resume/resume_generator.py
==========================

Generates tailored DOCX resumes for genuinely strong matches.

The resume generator itself does NOT invent experience, skills,
projects, metrics, or responsibilities.

All resume content comes from profile.json through resume_data.py.
"""

import json
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from app.config import (
    GENERATED_DIR,
    MASTER_RESUME,
)

from app.resume.resume_data import (
    load_full_profile,
    select_relevant_content,
)


# ============================================================
# STYLING
# ============================================================

def _set_font(
    run,
    name: str = "Calibri",
    size: int = 11,
    bold: bool = False,
    italic: bool = False,
    color: tuple = None,
):
    """Apply font formatting."""

    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic

    if color:
        run.font.color.rgb = RGBColor(*color)


def _add_horizontal_rule(doc: Document):
    """Add a thin horizontal divider."""

    p = doc.add_paragraph()

    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)

    pPr = p._p.get_or_add_pPr()

    pBdr = OxmlElement("w:pBdr")

    bottom = OxmlElement("w:bottom")

    bottom.set(
        qn("w:val"),
        "single"
    )

    bottom.set(
        qn("w:sz"),
        "6"
    )

    bottom.set(
        qn("w:space"),
        "1"
    )

    bottom.set(
        qn("w:color"),
        "2E4057"
    )

    pBdr.append(bottom)

    pPr.append(pBdr)

    return p


def _section_heading(
    doc: Document,
    title: str,
):
    """Add section heading."""

    _add_horizontal_rule(doc)

    p = doc.add_paragraph()

    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)

    run = p.add_run(
        title.upper()
    )

    _set_font(
        run,
        size=11,
        bold=True,
        color=(46, 64, 87),
    )

    return p


def _bullet(
    doc: Document,
    text: str,
):
    """Add bullet paragraph."""

    if not text:
        return None

    p = doc.add_paragraph(
        style="List Bullet"
    )

    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)

    run = p.add_run(
        str(text)
    )

    _set_font(
        run,
        size=10,
    )

    return p


# ============================================================
# BUILD RESUME
# ============================================================

def build_resume_docx(
    company_name: str,
    job_title: str,
    selected_content: dict,
    output_path: Path,
) -> Path:
    """
    Build DOCX resume from selected profile content.
    """

    profile = load_full_profile()

    personal = profile.get(
        "personal",
        {}
    )

    # --------------------------------------------------------
    # Education
    # --------------------------------------------------------

    edu_raw = profile.get(
        "education",
        {}
    )

    if (
        isinstance(edu_raw, list)
        and edu_raw
    ):
        education = edu_raw[0]
    else:
        education = edu_raw

    # --------------------------------------------------------
    # Document
    # --------------------------------------------------------

    if MASTER_RESUME.exists():

        doc = Document(
            str(MASTER_RESUME)
        )

        # Remove existing body content.
        for element in list(
            doc.element.body
        ):
            doc.element.body.remove(
                element
            )

    else:

        doc = Document()

    # --------------------------------------------------------
    # Margins
    # --------------------------------------------------------

    section = doc.sections[0]

    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    # --------------------------------------------------------
    # Name
    # --------------------------------------------------------

    name = personal.get(
        "name",
        "Candidate"
    )

    p_name = doc.add_paragraph()

    p_name.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    p_name.paragraph_format.space_after = Pt(2)

    run = p_name.add_run(
        name
    )

    _set_font(
        run,
        size=20,
        bold=True,
        color=(46, 64, 87),
    )

    # --------------------------------------------------------
    # Contact
    # --------------------------------------------------------

    contact_parts = []

    for key in (
        "email",
        "phone",
        "linkedin",
        "github",
        "location",
    ):

        value = personal.get(
            key
        )

        if value:
            contact_parts.append(
                str(value)
            )

    if contact_parts:

        p_contact = doc.add_paragraph()

        p_contact.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        p_contact.paragraph_format.space_after = Pt(4)

        run = p_contact.add_run(
            "  |  ".join(contact_parts)
        )

        _set_font(
            run,
            size=9,
            color=(80, 80, 80),
        )

    # --------------------------------------------------------
    # Summary
    # --------------------------------------------------------

    summary_raw = selected_content.get(
        "summary",
        ""
    )

    if isinstance(summary_raw, dict):
        summary = summary_raw.get(
            "text",
            ""
        )
    else:
        summary = summary_raw

    if summary:

        _section_heading(
            doc,
            "Professional Summary"
        )

        p = doc.add_paragraph()

        p.paragraph_format.space_after = Pt(2)

        run = p.add_run(
            str(summary)
        )

        _set_font(
            run,
            size=10,
        )

    # --------------------------------------------------------
    # Education
    # --------------------------------------------------------

    if (
        isinstance(education, dict)
        and education.get("degree")
    ):

        _section_heading(
            doc,
            "Education"
        )

        p = doc.add_paragraph()

        p.paragraph_format.space_after = Pt(0)

        degree = education.get(
            "degree",
            ""
        )

        institution = education.get(
            "institution",
            ""
        )

        year = education.get(
            "graduation",
            education.get(
                "graduation_year",
                ""
            )
        )

        cgpa = education.get(
            "cgpa"
        )

        r_degree = p.add_run(
            str(degree)
        )

        _set_font(
            r_degree,
            size=10,
            bold=True,
        )

        if institution:

            r_inst = p.add_run(
                f"  |  {institution}"
            )

            _set_font(
                r_inst,
                size=10,
            )

        p2 = doc.add_paragraph()

        p2.paragraph_format.space_after = Pt(2)

        details = []

        if year:
            details.append(
                f"Graduating {year}"
            )

        if cgpa:
            details.append(
                f"CGPA: {cgpa}/10"
            )

        if details:

            r2 = p2.add_run(
                "  ".join(details)
            )

            _set_font(
                r2,
                size=9,
                color=(90, 90, 90),
            )

    # --------------------------------------------------------
    # Skills
    # --------------------------------------------------------

    skills = selected_content.get(
        "skills",
        []
    )

    if skills:

        _section_heading(
            doc,
            "Skills"
        )

        p = doc.add_paragraph()

        p.paragraph_format.space_after = Pt(2)

        run = p.add_run(
            ", ".join(
                str(skill)
                for skill in skills
            )
        )

        _set_font(
            run,
            size=10,
        )

    # --------------------------------------------------------
    # Experience
    # --------------------------------------------------------

    experience = [
        exp
        for exp in selected_content.get(
            "experience",
            []
        )
        if (
            isinstance(exp, dict)
            and exp.get("role")
            and exp.get("company")
        )
    ]

    if experience:

        _section_heading(
            doc,
            "Experience"
        )

        for exp in experience:

            p = doc.add_paragraph()

            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(0)

            r_role = p.add_run(
                exp.get(
                    "role",
                    ""
                )
            )

            _set_font(
                r_role,
                size=10,
                bold=True,
            )

            company = exp.get(
                "company",
                ""
            )

            start = exp.get(
                "start",
                ""
            )

            end = exp.get(
                "end",
                "Present"
            )

            date_range = ""

            if start or end:
                date_range = (
                    f"  {start} – {end}"
                )

            if company:

                r_company = p.add_run(
                    f"  at {company}"
                    f"{date_range}"
                )

                _set_font(
                    r_company,
                    size=10,
                    color=(90, 90, 90),
                )

            bullets = exp.get(
                "bullets",
                []
            )

            if isinstance(
                bullets,
                list
            ):

                for bullet in bullets:

                    if bullet:
                        _bullet(
                            doc,
                            bullet
                        )

    # --------------------------------------------------------
    # Projects
    # --------------------------------------------------------

    projects = [
        project
        for project in selected_content.get(
            "projects",
            []
        )
        if (
            isinstance(project, dict)
            and project.get("name")
        )
    ]

    if projects:

        _section_heading(
            doc,
            "Projects"
        )

        for project in projects:

            p = doc.add_paragraph()

            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(0)

            r_name = p.add_run(
                project.get(
                    "name",
                    ""
                )
            )

            _set_font(
                r_name,
                size=10,
                bold=True,
            )

            tech = project.get(
                "tech",
                []
            )

            if tech:

                r_tech = p.add_run(
                    "  |  "
                    + ", ".join(
                        str(t)
                        for t in tech
                    )
                )

                _set_font(
                    r_tech,
                    size=9,
                    color=(90, 90, 90),
                )

            description = project.get(
                "description",
                ""
            )

            if description:
                _bullet(
                    doc,
                    description
                )

            bullets = project.get(
                "bullets",
                []
            )

            if isinstance(
                bullets,
                list
            ):

                for bullet in bullets:

                    if bullet:
                        _bullet(
                            doc,
                            bullet
                        )

    # --------------------------------------------------------
    # Certifications
    # --------------------------------------------------------

    certifications = selected_content.get(
        "certifications",
        []
    )

    if certifications:

        _section_heading(
            doc,
            "Certifications"
        )

        for cert in certifications:

            if isinstance(
                cert,
                str
            ):
                text = cert

            elif isinstance(
                cert,
                dict
            ):
                text = cert.get(
                    "name",
                    ""
                )

            else:
                continue

            if text:
                _bullet(
                    doc,
                    text
                )

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    doc.save(
        str(output_path)
    )

    print(
        f"  [Resume] Saved: "
        f"{output_path.name}"
    )

    return output_path


# ============================================================
# SINGLE RESUME
# ============================================================

def generate_resume(
    company_name: str,
    job_title: str,
    job_description: str,
    matched_skills: list,
    keywords: list,
) -> Path:
    """
    Generate one tailored resume.
    """

    safe_company = re.sub(
        r"[^\w]+",
        "_",
        company_name.lower()
    ).strip("_")

    safe_title = re.sub(
        r"[^\w]+",
        "_",
        job_title.lower()
    ).strip("_")[:40]

    filename = (
        f"{safe_company}_{safe_title}.docx"
    )

    output_path = (
        GENERATED_DIR / filename
    )

    selected = select_relevant_content(
        job_title=job_title,
        job_description=job_description,
        matched_skills=matched_skills,
        keywords=keywords,
    )

    return build_resume_docx(
        company_name=company_name,
        job_title=job_title,
        selected_content=selected,
        output_path=output_path,
    )


# ============================================================
# STRONG MATCH RESUMES
# ============================================================

def generate_resumes_for_strong_matches() -> list[Path]:
    """
    Generate resumes only for strong, technically relevant roles.

    IMPORTANT:
        This function assumes analysis has already passed the
        matching/AI relevance gates.
    """

    from app.database.db import get_connection

    generated = []

    with get_connection() as conn:

        rows = conn.execute(
            """
            SELECT
                j.id,
                j.company,
                j.title,
                j.description,
                a.matched_skills,
                a.keywords,
                a.recommendation,
                a.final_score,
                a.role_type
            FROM jobs j
            JOIN analysis a
                ON j.id = a.job_id
            WHERE
                a.recommendation = 'strong'
                AND a.final_score >= 0.70
                AND a.role_type IS NOT NULL
                AND a.role_type != 'Other'
            ORDER BY
                a.final_score DESC
            """
        ).fetchall()

    print(
        f"[Resume] Strong matches eligible: "
        f"{len(rows)}"
    )

    for row in rows:

        try:

            matched_skills = json.loads(
                row["matched_skills"]
                or "[]"
            )

            keywords = json.loads(
                row["keywords"]
                or "[]"
            )

            path = generate_resume(
                company_name=row["company"],
                job_title=row["title"],
                job_description=(
                    row["description"]
                    or ""
                ),
                matched_skills=matched_skills,
                keywords=keywords,
            )

            generated.append(path)

        except Exception as e:

            print(
                f"[Resume] Failed for "
                f"{row['company']} | "
                f"{row['title']}: "
                f"{type(e).__name__}: {e}"
            )

    print(
        f"[Resume] Generated "
        f"{len(generated)} resumes."
    )

    return generated